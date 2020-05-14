#importing files

from docx import Document
from fpdf import FPDF
import os

#initialization...
print("Welcome to the file converter...")
path_file = input("Enter the path of your file")

print("To convert your txt file to pdf choose 1")
print("To convert your txt file to docx choose 2")
i = int(input())

if(i == 1):
	pdf = FPDF()
	pdf.add_page()

	#Set the font size...
	#The Options for font style are,Arial,Courier,Times(arial in default)

	d = int(input("Enter the font size..."))
	pdf.set_font("Arial", size=d)

	# open the text file in read mode
	f = open(path_file, "r")

	# insert the texts in pdf
	for x in f:
		pdf.cell(0, 10, txt=x, ln=1, align='L')
	'''fpdf.cell(w, h = 0, txt = '', border = 0, ln = 0, align = '')
		w is cell width
		h is cell height
		align can be L,C,R(left,centre,right align)'''


	# save the pdf
	name = input("Enter the output pdf name...")
	pdf.output(name + ".pdf")

elif(i == 2):
	doc = Document()

	with open(path_file, 'r', encoding='utf-8') as file:
		doc.add_paragraph(file.read())

	name = input("Enter the output docx name...")
	doc.save(name + ".docx")
	os.startfile(name + ".docx")

else:
	print("Invalid Choice")



